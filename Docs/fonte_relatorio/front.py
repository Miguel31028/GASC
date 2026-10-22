from docx.shared import Cm

from lib import set_par_text

TITULO = 'Barcelos Solidário: Análise e Modelação do Sistema de Gestão do Espaço Social do GASC'

RESUMO = [
    'As instituições de solidariedade social que distribuem bens alimentares lidam com doações irregulares, em formatos e quantidades muito variados e com prazos de validade por vezes curtos. Este trabalho, desenvolvido no âmbito do projeto Barcelos Solidário (GASC, Câmara Municipal de Barcelos e IPCA), tem como objetivo conceber um sistema de informação para gerir o Espaço Social do GASC – Grupo de Acção Social Cristã: a receção de bens, o armazenamento em prateleiras e equipamentos de frio, o controlo de validades, a distribuição de cabazes e outros bens aos agregados apoiados e a gestão das refeições.',
    'O documento apresenta a análise do processo atual, o levantamento e a priorização de requisitos funcionais e não funcionais pelo método MoSCoW, as regras de negócio e a modelação completa do sistema em UML (casos de uso, atividades, classes, estados e sequências), dividida por áreas para preservar a legibilidade. Inclui o modelo entidade-relacionamento, o dicionário de dados, o modelo em Cloud Firestore e 42 protótipos de alta fidelidade. A arquitetura assenta em aplicações Android em Kotlin e no Firebase, com as operações críticas em Cloud Functions transacionais e segurança aplicada no servidor. Destacam-se a sugestão de localização, a seleção FEFO, a congelação de produtos em fim de prazo, a identificação por credencial QR e a geração automática do registo HACCP RGDHSA049.',
    'O âmbito foi dividido entre um produto mínimo viável, focado na operação diária, e um conjunto de melhorias futuras documentadas, como o reconhecimento de produtos e validades por imagem, a leitura automática de guias, a identificação biométrica, o cabaz automático equilibrado e a sugestão de pratos. O resultado é uma especificação detalhada, com rastreabilidade entre requisitos, casos de uso, protótipos e testes, a validar com a instituição antes da implementação.',
]
KEYWORDS = '**Palavras-chave:** gestão de stock; FEFO; desperdício alimentar; análise de sistemas; UML; Kotlin; Firebase; IPSS.'

ABSTRACT = [
    'Social solidarity institutions that distribute food deal with irregular donations, in widely varying formats and quantities and often with short expiry dates. This work, carried out within the Barcelos Solidário project (GASC, Barcelos City Council and IPCA), aims to design an information system to manage the Social Space of GASC – Grupo de Acção Social Cristã: the reception of goods, storage on shelves and in cold equipment, expiry date control, the distribution of food baskets and other goods to the supported households, and meal management.',
    'The document presents the analysis of the current process, the elicitation and MoSCoW prioritisation of functional and non-functional requirements, the business rules and a complete UML model of the system (use cases, activities, classes, states and sequences), split by area to preserve readability. It includes the entity-relationship model, the data dictionary, the Cloud Firestore model and 42 high-fidelity prototypes. The architecture is based on Android applications written in Kotlin and on Firebase, with critical operations in transactional Cloud Functions and server-side security. Key features include storage location suggestions, FEFO selection, freezing of products close to expiry, identification through QR credentials and automatic generation of the HACCP record RGDHSA049.',
    'The scope was divided into a minimum viable product, focused on daily operations, and a set of documented future improvements, such as image-based recognition of products and expiry dates, automatic reading of delivery notes, biometric identification, automatically balanced food baskets and dish suggestions. The result is a detailed specification, with traceability between requirements, use cases, prototypes and tests, to be validated with the institution before implementation.',
]
ABS_KEYWORDS = '**Keywords:** inventory management; FEFO; food waste; systems analysis; UML; Kotlin; Firebase; non-profit organisations.'

GLOSSARIO = [
    ('Agendamento', 'Marcação de uma data e hora para a entrega de um cabaz ou de outros bens a um agregado.'),
    ('Agregado familiar', 'Conjunto de pessoas que vivem em economia comum e são apoiadas em conjunto; é a unidade a que se atribuem cabazes e pedidos de refeição.'),
    ('Apresentação', 'Forma de embalagem em que um produto é recebido (unidade, caixa, palete, saco), com um fator de conversão para a unidade de stock.'),
    ('Cabaz', 'Conjunto de bens alimentares entregue a um agregado familiar numa data agendada.'),
    ('Callable', 'Cloud Function invocada diretamente pela aplicação, com o utilizador autenticado, para executar uma operação no servidor.'),
    ('Célula', 'Posição de uma prateleira identificada pelo código xxx-xx-xx (coluna, linha e andar).'),
    ('Cloud Function', 'Função executada no servidor da Google em resposta a um pedido ou a um evento, sem gestão de servidores.'),
    ('Credencial QR', 'Código QR pessoal, em papel ou na App Utente, que identifica uma pessoa autorizada a levantar bens.'),
    ("Crow's foot", 'Notação de diagramas entidade-relacionamento em que as cardinalidades são representadas por símbolos nas extremidades das ligações.'),
    ('Custom claims', 'Atributos adicionados ao *token* de autenticação do Firebase, usados para representar os perfis de utilizador.'),
    ('Desnormalização', 'Repetição controlada de dados noutro documento para evitar leituras adicionais (ex.: nome do produto guardado no lote).'),
    ('Desperdício', 'Quantidade de produto retirada do stock por o produto ter expirado ou estar danificado ou não conforme.'),
    ('Guia de transporte', 'Documento que acompanha os bens entregues por uma empresa, com a identificação dos artigos e das quantidades.'),
    ('Idempotência', 'Propriedade de uma operação que, repetida, produz o mesmo resultado que uma única execução.'),
    ('Levantamento', 'Registo da entrega de refeições a uma pessoa autorizada de um agregado.'),
    ('Lote', 'Conjunto de unidades do mesmo produto (nome, marca e formato) com a mesma validade, recebidas na mesma receção, gerido como uma unidade de stock.'),
    ('Mockup', 'Protótipo de alta fidelidade de um ecrã, que define o conteúdo, os componentes e o comportamento esperado.'),
    ('MoSCoW', "Método de priorização de requisitos em *Must*, *Should*, *Could* e *Won't*."),
    ('Pedido de refeição', 'Intenção registada por um agregado de receber um número de refeições num determinado dia.'),
    ('Produto a granel', 'Produto cuja unidade de stock é o peso ou o volume (ex.: batatas em kg), contado também em volumes (sacos, caixas).'),
    ('Security Rules', 'Regras declarativas do Firebase que controlam o acesso de leitura e escrita aos dados.'),
    ('Serviço de refeição', 'Refeição (almoço ou jantar) servida num determinado dia.'),
    ('Subcoleção', 'Coleção do Firestore associada a um documento (ex.: as pessoas de um agregado).'),
    ('Swimlane', 'Raia de um diagrama de atividade que agrupa as ações de um mesmo ator ou componente.'),
    ('Tabela de congelação', 'Tabela, definida por categoria, com os dias de extensão da validade propostos quando um produto é congelado.'),
    ('Token', 'Valor aleatório ou credencial digital que identifica um utilizador ou uma pessoa sem revelar dados pessoais.'),
    ('Trigger', 'Cloud Function executada automaticamente quando ocorre um evento, como a criação de um documento.'),
    ('Utente', 'Pessoa apoiada pela instituição, membro de um agregado familiar registado.'),
    ('Volume', 'Unidade física de um produto a granel (saco, caixa) cujo peso pode ser consumido parcialmente.'),
    ('Zona de frio', 'Divisão de um equipamento de frio identificada pelo lado (D, C, E) e pelo nível.'),
]

SIGLAS = [
    ('ADM, RESP, TEC, OP, COZ, UT', 'Perfis Administrador, Responsável da instituição, Técnico de ação social, Operador de armazém, Cozinha e Utente (matriz de permissões)'),
    ('AIPD', 'Avaliação de Impacto sobre a Proteção de Dados'),
    ('API', 'Application Programming Interface (Interface de Programação de Aplicações)'),
    ('BaaS', 'Backend as a Service (Backend como Serviço)'),
    ('C / NC', 'Conforme / Não conforme'),
    ('CI', 'Continuous Integration (Integração Contínua)'),
    ('CMB', 'Câmara Municipal de Barcelos'),
    ('CSV', 'Comma-Separated Values (valores separados por vírgulas)'),
    ('dp', 'Density-independent pixel (pixel independente da densidade do ecrã)'),
    ('EAN', 'European Article Number (código de barras de artigos)'),
    ('ER', 'Entidade-Relacionamento'),
    ('ERP', 'Enterprise Resource Planning (Planeamento de Recursos Empresariais)'),
    ('FAO', 'Food and Agriculture Organization (Organização das Nações Unidas para a Alimentação e a Agricultura)'),
    ('FCM', 'Firebase Cloud Messaging'),
    ('FEFO', 'First Expired, First Out (primeiro a expirar, primeiro a sair)'),
    ('FIFO', 'First In, First Out (primeiro a entrar, primeiro a sair)'),
    ('GASC', 'Grupo de Acção Social Cristã'),
    ('HACCP', 'Hazard Analysis and Critical Control Points (Análise de Perigos e Pontos Críticos de Controlo)'),
    ('HTTPS', 'HyperText Transfer Protocol Secure'),
    ('IA', 'Inteligência Artificial'),
    ('IoT', 'Internet of Things (Internet das Coisas)'),
    ('IPCA', 'Instituto Politécnico do Cávado e do Ave'),
    ('IPSS', 'Instituição Particular de Solidariedade Social'),
    ('ISO/IEC', 'International Organization for Standardization / International Electrotechnical Commission'),
    ('JSON', 'JavaScript Object Notation'),
    ('LLM', 'Large Language Model (Grande Modelo de Linguagem)'),
    ('MVP', 'Minimum Viable Product (Produto Mínimo Viável)'),
    ('MVVM', 'Model-View-ViewModel'),
    ('NIF', 'Número de Identificação Fiscal'),
    ('NISS', 'Número de Identificação da Segurança Social'),
    ('NoSQL', 'Not only SQL (base de dados não relacional)'),
    ('OCR', 'Optical Character Recognition (Reconhecimento Ótico de Caracteres)'),
    ('PC', 'Personal Computer (computador pessoal)'),
    ('PDF', 'Portable Document Format'),
    ('QR', 'Quick Response (código de resposta rápida)'),
    ('RF', 'Requisito Funcional'),
    ('RGPD', 'Regulamento Geral sobre a Proteção de Dados'),
    ('RN', 'Regra de Negócio'),
    ('RNF', 'Requisito Não Funcional'),
    ('SDK', 'Software Development Kit (kit de desenvolvimento de software)'),
    ('SMS', 'Short Message Service (mensagem de texto)'),
    ('SUS', 'System Usability Scale (escala de usabilidade de sistemas)'),
    ('TA', 'Teste de Aceitação'),
    ('TLS', 'Transport Layer Security'),
    ('UC', 'Use Case (Caso de Uso)'),
    ('UI', 'User Interface (Interface com o Utilizador)'),
    ('UID', 'User Identifier (identificador do utilizador no Firebase Authentication)'),
    ('UML', 'Unified Modeling Language (Linguagem de Modelação Unificada)'),
    ('WCAG', 'Web Content Accessibility Guidelines'),
    ('WMS', 'Warehouse Management System (Sistema de Gestão de Armazém)'),
]


def _remove(p):
    p._p.getparent().remove(p._p)


def _is_plain_empty(p):
    x = p._p.xml
    return not p.text.strip() and 'sectPr' not in x and 'w:br' not in x and 'fldChar' not in x and 'drawing' not in x


def _insert_before(anchor_el, r, text, style='Normal', indent=False):
    p = r.doc.add_paragraph(style=style)
    r.add_runs(p, text)
    if not indent:
        p.paragraph_format.first_line_indent = Cm(0)
    anchor_el.addprevious(p._p)
    return p


def front(r):
    """Preenche as páginas preliminares. Devolve as âncoras para os índices de figuras e tabelas."""
    P = list(r.doc.paragraphs)
    set_par_text(P[5], '')
    set_par_text(P[8], TITULO)
    set_par_text(P[17], 'Ano letivo 2026/2027')
    set_par_text(P[45], 'Informação sobre o Estágio')
    set_par_text(P[46], 'GASC – Grupo de Acção Social Cristã')
    set_par_text(P[47], 'Campo 25 de Abril, n.º 24, 4750-102 Barcelos')

    # Resumo
    set_par_text(P[57], RESUMO[0])
    set_par_text(P[58], RESUMO[1])
    set_par_text(P[59], RESUMO[2])
    kw = _insert_before(P[60]._p, r, KEYWORDS)
    kw.paragraph_format.space_before = __import__('docx').shared.Pt(12)

    # Abstract
    set_par_text(P[67], 'ABSTRACT')
    set_par_text(P[69], ABSTRACT[0])
    for t in ABSTRACT[1:]:
        p = _insert_before(P[70]._p, r, t, indent=True)
    _insert_before(P[70]._p, r, ABS_KEYWORDS).paragraph_format.space_before = __import__('docx').shared.Pt(12)
    for i in range(70, 74):
        if _is_plain_empty(P[i]):
            _remove(P[i])

    # Agradecimentos
    set_par_text(P[81], '[Texto de agradecimentos a redigir pelo autor.]')
    set_par_text(P[95], '')

    # Índice
    set_par_text(P[103], '')

    # Índice de figuras: remove o campo do modelo e os espaços em branco
    fig_anchor = P[113]._p
    _remove(P[111]); _remove(P[112])
    for i in range(114, 131):
        if _is_plain_empty(P[i]) or i == 130:
            _remove(P[i])

    # Índice de tabelas
    tab_anchor = P[139]._p
    _remove(P[136]); _remove(P[137]); _remove(P[138])
    for i in range(140, 157):
        if _is_plain_empty(P[i]) or i == 156:
            _remove(P[i])

    # Glossário
    anchor = P[164]._p
    for term, d in GLOSSARIO:
        _insert_before(anchor, r, f'**{term}** – {d}')
    for i in (161, 162, 163):
        _remove(P[i])
    for i in range(164, 180):
        if _is_plain_empty(P[i]) or i == 179:
            _remove(P[i])

    # Siglas e acrónimos
    anchor = P[189]._p
    for s, d in SIGLAS:
        q = _insert_before(anchor, r, f'**{s}** – {d}')
        q.paragraph_format.space_after = __import__('docx').shared.Pt(0)
    for i in (185, 186, 187, 188):
        _remove(P[i])
    for i in range(189, 205):
        if _is_plain_empty(P[i]) or i == 204:
            _remove(P[i])
    return fig_anchor, tab_anchor
