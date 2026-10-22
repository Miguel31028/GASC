"""Motor de construção do relatório sobre o modelo Word do IPCA."""
import copy
import re

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt

TEXT_WIDTH_DXA = 9071  # A4 (11906) - margens 1701 + 1134
MAX_IMG_W_CM = 15.5
MAX_IMG_H_CM = 21.0


def _q(tag):
    return qn(tag)


TBLPR_ORDER = ['tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize', 'tblStyleColBandSize', 'tblW', 'jc',
               'tblCellSpacing', 'tblInd', 'tblBorders', 'shd', 'tblLayout', 'tblCellMar', 'tblLook', 'tblCaption', 'tblDescription']


def set_tblpr(tblpr, xml_children):
    """Substitui/insere filhos de w:tblPr respeitando a ordem do esquema."""
    for xml in xml_children:
        el = parse_xml(xml)
        name = el.tag.split('}')[1]
        for old in tblpr.findall(_q('w:' + name)):
            tblpr.remove(old)
        idx = TBLPR_ORDER.index(name)
        after = [c for c in tblpr if c.tag.split('}')[1] in TBLPR_ORDER and TBLPR_ORDER.index(c.tag.split('}')[1]) > idx]
        if after:
            after[0].addprevious(el)
        else:
            tblpr.append(el)


def set_tcw(cell, w):
    tcpr = cell._tc.get_or_add_tcPr()
    for old in tcpr.findall(_q('w:tcW')):
        tcpr.remove(old)
    tcpr.insert(0, parse_xml(f'<w:tcW {nsdecls("w")} w:w="{w}" w:type="dxa"/>'))
    return tcpr


class Report:
    def __init__(self, template, pages=None):
        self.doc = docx.Document(template)
        self.body = self.doc.element.body
        self.pages = pages or {}
        self.bm_id = 1000
        self.hnum = [0, 0, 0, 0, 0]
        self.headings = []  # (level, number, text, bookmark)
        self.figs = []  # (number, text, bookmark)
        self.tabs = []
        self.chapter_count = 0
        self.list_num_ids = {}
        self._setup_numbering()
        self._collect_sections()

    # ------------------------------------------------------------------
    # preparação do modelo
    def _collect_sections(self):
        sect_paras = [p for p in self.body.iterchildren(_q('w:p'))
                      if p.find(_q('w:pPr')) is not None and p.find(_q('w:pPr')).find(_q('w:sectPr')) is not None]
        self.sect = [copy.deepcopy(p.find(_q('w:pPr')).find(_q('w:sectPr'))) for p in sect_paras]
        # o modelo reinicia os anexos em numeração romana (XIX...), que repete a das páginas
        # preliminares; continua-se a numeração árabe do corpo do documento
        for sp in (self.sect[11], self.body.find(_q('w:sectPr'))):
            for el in sp.findall(_q('w:pgNumType')):
                sp.remove(el)
        self.front_end = sect_paras[5]  # fim das páginas preliminares

    def strip_template_body(self):
        """Remove os capítulos de exemplo do modelo (tudo após as páginas preliminares)."""
        children = list(self.body)
        start = children.index(self.front_end) + 1
        for el in children[start:]:
            if el.tag == _q('w:sectPr'):
                continue
            self.body.remove(el)

    def _setup_numbering(self):
        numbering = self.doc.part.numbering_part.element
        first_num = numbering.find(_q('w:num'))
        bullet = parse_xml(
            f'<w:abstractNum {nsdecls("w")} w:abstractNumId="90"><w:multiLevelType w:val="hybridMultilevel"/>'
            '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="•"/>'
            '<w:lvlJc w:val="left"/><w:pPr><w:ind w:left="1069" w:hanging="360"/></w:pPr>'
            '<w:rPr><w:rFonts w:ascii="Arial" w:hAnsi="Arial"/></w:rPr></w:lvl>'
            '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="–"/>'
            '<w:lvlJc w:val="left"/><w:pPr><w:ind w:left="1789" w:hanging="360"/></w:pPr></w:lvl>'
            '</w:abstractNum>')
        decimal = parse_xml(
            f'<w:abstractNum {nsdecls("w")} w:abstractNumId="91"><w:multiLevelType w:val="hybridMultilevel"/>'
            '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/>'
            '<w:lvlJc w:val="left"/><w:pPr><w:ind w:left="1069" w:hanging="360"/></w:pPr></w:lvl>'
            '</w:abstractNum>')
        first_num.addprevious(bullet)
        first_num.addprevious(decimal)
        numbering.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="90"><w:abstractNumId w:val="90"/></w:num>'))
        self.numbering = numbering
        self.next_num = 91

    def _new_decimal_num(self):
        nid = self.next_num
        self.next_num += 1
        self.numbering.append(parse_xml(
            f'<w:num {nsdecls("w")} w:numId="{nid}"><w:abstractNumId w:val="91"/>'
            '<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride></w:num>'))
        return nid

    # ------------------------------------------------------------------
    # texto com marcação simples: **negrito**, *itálico*, `código`
    TOKEN = re.compile(r'(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)')

    def add_runs(self, p, text, size=None, bold=False):
        for part in self.TOKEN.split(text):
            if not part:
                continue
            if part.startswith('**'):
                r = p.add_run(part[2:-2]); r.bold = True
            elif part.startswith('`'):
                r = p.add_run(part[1:-1]); r.font.name = 'Courier New'
            elif part.startswith('*'):
                r = p.add_run(part[1:-1]); r.italic = True
            else:
                r = p.add_run(part)
            if bold:
                r.bold = True
            if size:
                r.font.size = Pt(size)
        return p

    def _bookmark(self, p, name):
        self.bm_id += 1
        start = parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="{self.bm_id}" w:name="{name}"/>')
        end = parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="{self.bm_id}"/>')
        ppr = p._p.find(_q('w:pPr'))
        if ppr is not None:
            ppr.addnext(start)
        else:
            p._p.insert(0, start)
        p._p.append(end)

    # ------------------------------------------------------------------
    # blocos
    def para(self, text, style='Normal', align=None, indent=True, keep_next=False):
        p = self.doc.add_paragraph(style=style)
        self.add_runs(p, text)
        if align == 'center':
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if not indent:
            p.paragraph_format.first_line_indent = Cm(0)
        if keep_next:
            p.paragraph_format.keep_with_next = True
        return p

    def paras(self, *texts):
        for t in texts:
            self.para(t)

    def empty(self, n=1):
        for _ in range(n):
            self.doc.add_paragraph()

    def bullets(self, items, level=0):
        for it in items:
            if isinstance(it, (list, tuple)):
                self.bullets(it, level + 1)
                continue
            p = self.doc.add_paragraph(style='List Paragraph')
            self._numpr(p, 90, level)
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_after = Pt(2)
            self.add_runs(p, it)

    def numbered(self, items):
        nid = self._new_decimal_num()
        for it in items:
            p = self.doc.add_paragraph(style='List Paragraph')
            self._numpr(p, nid, 0)
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_after = Pt(2)
            self.add_runs(p, it)

    def _numpr(self, p, num_id, level):
        ppr = p._p.get_or_add_pPr()
        numpr = parse_xml(f'<w:numPr {nsdecls("w")}><w:ilvl w:val="{level}"/><w:numId w:val="{num_id}"/></w:numPr>')
        ppr.insert(1 if ppr.find(_q('w:pStyle')) is not None else 0, numpr)

    def code(self, lines):
        for ln in lines.strip('\n').split('\n'):
            p = self.doc.add_paragraph(style='Code')
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.left_indent = Cm(0.5)
            r = p.add_run(ln if ln else ' ')
            r.font.size = Pt(9)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    # ------------------------------------------------------------------
    # títulos e secções
    def _section_break(self, sect):
        p = self.doc.add_paragraph()
        p._p.get_or_add_pPr().append(copy.deepcopy(sect))

    def chapter(self, title):
        self.chapter_count += 1
        if self.chapter_count == 2:
            self._section_break(self.sect[6])  # fecha o capítulo 1 (numeração árabe a partir de 1)
        elif self.chapter_count > 2:
            self._section_break(self.sect[7])  # quebra para página ímpar
        self.empty(4)
        self.heading(title, 1)

    def end_chapters(self):
        self._section_break(self.sect[7] if self.chapter_count > 1 else self.sect[6])

    def heading(self, title, level):
        self.hnum[level - 1] += 1
        for i in range(level, 5):
            self.hnum[i] = 0
        number = '.'.join(str(n) for n in self.hnum[:level]) + '.'
        p = self.doc.add_paragraph(style=f'Heading {level}')
        p.add_run(title)
        bm = f'_Toc9{len(self.headings):05d}'
        self._bookmark(p, bm)
        self.headings.append((level, number, title, bm))
        return p

    def h2(self, t):
        return self.heading(t, 2)

    def h3(self, t):
        return self.heading(t, 3)

    def unnumbered_title(self, text, toc=False):
        p = self.doc.add_paragraph(style='Titulo Não Numerado')
        p.add_run(text)
        if toc:
            bm = f'_Toc9{len(self.headings):05d}'
            self._bookmark(p, bm)
            self.headings.append((1, '', text, bm))
        return p

    def page_break(self):
        p = self.doc.add_paragraph()
        p.add_run().add_break(docx.enum.text.WD_BREAK.PAGE)

    # ------------------------------------------------------------------
    # figuras e tabelas
    def _caption(self, kind, text, store):
        n = len(store) + 1
        bm = f'_Toc8{kind[0]}{n:04d}'
        p = self.doc.add_paragraph(style='Caption')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(10)
        p.add_run(f'{kind} ')
        fld = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr=" SEQ {kind} \\* ARABIC "><w:r><w:rPr><w:noProof/></w:rPr><w:t>{n}</w:t></w:r></w:fldSimple>')
        p._p.append(fld)
        p.add_run(f': {text}')
        self._bookmark(p, bm)
        store.append((n, text, bm))
        return n

    def next_fig(self):
        return len(self.figs) + 1

    def next_tab(self):
        return len(self.tabs) + 1

    def tab_no(self, prefix):
        return next(n for n, t, _ in self.tabs if t.startswith(prefix))

    def fig_no(self, prefix):
        return next(n for n, t, _ in self.figs if t.startswith(prefix))

    def figure(self, path, caption, max_w=MAX_IMG_W_CM, max_h=MAX_IMG_H_CM):
        from PIL import Image
        with Image.open(path) as im:
            w, h = im.size
        width = max_w
        if width * h / w > max_h:
            width = max_h * w / h
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = p.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.keep_with_next = True
        pf.space_before = Pt(6)
        pf.line_spacing = 1.0
        p.add_run().add_picture(path, width=Cm(width))
        return self._caption('Figura', caption, self.figs)

    def table(self, header, rows, caption, widths=None, size=9, first_bold=False):
        ncols = len(header)
        if widths is None:
            widths = [1] * ncols
        tot = sum(widths)
        dxa = [int(TEXT_WIDTH_DXA * w / tot) for w in widths]
        dxa[-1] = TEXT_WIDTH_DXA - sum(dxa[:-1])
        t = self.doc.add_table(rows=1 + len(rows), cols=ncols)
        t.style = self.doc.styles['Table Grid']
        set_tblpr(t._tbl.tblPr, [f'<w:tblW {nsdecls("w")} w:w="{TEXT_WIDTH_DXA}" w:type="dxa"/>',
                                 f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'])
        grid = t._tbl.tblGrid
        for gc, w in zip(grid.findall(_q('w:gridCol')), dxa):
            gc.set(_q('w:w'), str(w))
        for ri, row in enumerate([header] + list(rows)):
            tr = t.rows[ri]
            if ri == 0:
                trpr = tr._tr.get_or_add_trPr()
                trpr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            trpr = tr._tr.get_or_add_trPr()
            trpr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            for ci, val in enumerate(row):
                cell = tr.cells[ci]
                tcpr = set_tcw(cell, dxa[ci])
                if ri == 0:
                    tcpr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="D9D9D9"/>'))
                lines = str(val).split('\n')
                for li, ln in enumerate(lines):
                    p = cell.paragraphs[0] if li == 0 else cell.add_paragraph()
                    pf = p.paragraph_format
                    pf.first_line_indent = Cm(0)
                    pf.space_after = Pt(0)
                    pf.space_before = Pt(0)
                    pf.line_spacing = 1.0
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    if ln.startswith('- '):
                        ln = '• ' + ln[2:]
                    self.add_runs(p, ln, size=size, bold=(ri == 0 or (first_bold and ci == 0)))
        return self._caption('Tabela', caption, self.tabs)

    # ------------------------------------------------------------------
    # índices (TOC, figuras, tabelas) com números de página em cache
    def _entry_runs(self, anchor, parts, page):
        runs = ''
        for kind, val in parts:
            if kind == 'tab':
                runs += '<w:r><w:tab/></w:r>'
            else:
                runs += f'<w:r><w:rPr><w:noProof/></w:rPr><w:t xml:space="preserve">{_esc(val)}</w:t></w:r>'
        runs += ('<w:r><w:rPr><w:noProof/><w:webHidden/></w:rPr><w:tab/></w:r>'
                 '<w:r><w:rPr><w:noProof/><w:webHidden/></w:rPr><w:fldChar w:fldCharType="begin"/></w:r>'
                 f'<w:r><w:rPr><w:noProof/><w:webHidden/></w:rPr><w:instrText xml:space="preserve"> PAGEREF {anchor} \\h </w:instrText></w:r>'
                 '<w:r><w:rPr><w:noProof/><w:webHidden/></w:rPr><w:fldChar w:fldCharType="separate"/></w:r>'
                 f'<w:r><w:rPr><w:noProof/><w:webHidden/></w:rPr><w:t>{page}</w:t></w:r>'
                 '<w:r><w:rPr><w:noProof/><w:webHidden/></w:rPr><w:fldChar w:fldCharType="end"/></w:r>')
        return f'<w:hyperlink w:anchor="{anchor}" w:history="1">{runs}</w:hyperlink>'

    def _index_paragraphs(self, entries, instr, style_for, tabs_for):
        """entries: (style, anchor, parts). Devolve lista de elementos <w:p>."""
        out = []
        begin = ('<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
                 f'<w:r><w:instrText xml:space="preserve"> {instr} </w:instrText></w:r>'
                 '<w:r><w:fldChar w:fldCharType="separate"/></w:r>')
        for i, (lvl, anchor, parts) in enumerate(entries):
            page = self.pages.get(anchor, '')
            xml = (f'<w:p {nsdecls("w")}><w:pPr><w:pStyle w:val="{style_for(lvl)}"/>'
                   f'<w:tabs>{tabs_for(lvl)}</w:tabs><w:spacing w:after="40" w:line="276" w:lineRule="auto"/>'
                   '<w:ind w:firstLine="0"/><w:rPr><w:noProof/></w:rPr></w:pPr>'
                   + (begin if i == 0 else '') + self._entry_runs(anchor, parts, page) + '</w:p>')
            out.append(parse_xml(xml))
        out.append(parse_xml(f'<w:p {nsdecls("w")}><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'))
        return out

    def build_toc(self):
        sdt = self.body.find(_q('w:sdt'))
        content = sdt.find(_q('w:sdtContent'))
        first = content[0]
        for el in list(content):
            content.remove(el)
        content.append(first)
        left = {1: 0, 2: 220, 3: 440}
        numtab = {1: 600, 2: 1100, 3: 1700}
        entries = [(lvl, bm, [('t', num), ('tab', None), ('t', title)] if num else [('t', title)])
                   for lvl, num, title, bm in self.headings if lvl <= 3]
        for p in self._index_paragraphs(
                entries, 'TOC \\o "1-3" \\h \\z \\u', lambda l: f'ndice{l}',
                lambda l: f'<w:tab w:val="left" w:pos="{left[l] + numtab[l]}"/><w:tab w:val="right" w:leader="dot" w:pos="9061"/>'):
            content.append(p)

    def build_list_index(self, anchor_para, items, kind):
        """Substitui o parágrafo de campo do modelo pelo índice de figuras/tabelas."""
        entries = [(0, bm, [('t', f'{kind} {n}: {text}')]) for n, text, bm in items]
        new = self._index_paragraphs(entries, f'TOC \\h \\z \\c "{kind}"', lambda l: 'ndicedeilustraes',
                                     lambda l: '<w:tab w:val="right" w:leader="dot" w:pos="9061"/>')
        for el in new:
            anchor_para.addprevious(el)

    def enable_update_fields(self):
        settings = self.doc.settings.element
        el = parse_xml(f'<w:updateFields {nsdecls("w")} w:val="true"/>')
        later = ('hdrShapeDefaults', 'footnotePr', 'endnotePr', 'compat', 'docVars', 'rsids', 'mathPr', 'attachedSchema',
                 'themeFontLang', 'clrSchemeMapping', 'doNotIncludeSubdocsInStats', 'doNotAutoCompressPictures', 'forceUpgrade',
                 'captions', 'readModeInkLockDown', 'smartTagType', 'schemaLibrary', 'shapeDefaults', 'doNotEmbedSmartTags',
                 'decimalSymbol', 'listSeparator')
        nxt = next((c for c in settings if c.tag.split('}')[1] in later), None)
        if nxt is not None:
            nxt.addprevious(el)
        else:
            settings.append(el)

    def save(self, path):
        self.doc.save(path)


def _esc(s):
    return s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')


def set_par_text(p, text, keep_runs=1):
    """Substitui o texto de um parágrafo mantendo a formatação do primeiro run."""
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def figure_side(r, path, caption, paragraphs, img_w_cm=6.6, size=10):
    """Figura à esquerda e texto à direita, numa tabela sem margens; legenda por baixo."""
    left = int(TEXT_WIDTH_DXA * 0.5)
    right = TEXT_WIDTH_DXA - left
    t = r.doc.add_table(rows=1, cols=2)
    set_tblpr(t._tbl.tblPr, [
        f'<w:tblW {nsdecls("w")} w:w="{TEXT_WIDTH_DXA}" w:type="dxa"/>',
        f'<w:tblBorders {nsdecls("w")}>' + ''.join(f'<w:{b} w:val="nil"/>' for b in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV')) + '</w:tblBorders>',
        f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'])
    for gc, w in zip(t._tbl.tblGrid.findall(_q('w:gridCol')), (left, right)):
        gc.set(_q('w:w'), str(w))
    tr = t.rows[0]
    tr._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    c0, c1 = tr.cells
    for c, w in ((c0, left), (c1, right)):
        set_tcw(c, w)
    c1._tc.get_or_add_tcPr().append(parse_xml(f'<w:vAlign {nsdecls("w")} w:val="top"/>'))
    p = c0.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(path, width=Cm(img_w_cm))
    for i, txt in enumerate(paragraphs):
        q = c1.paragraphs[0] if i == 0 else c1.add_paragraph()
        pf = q.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.line_spacing = 1.15
        pf.space_after = Pt(6)
        q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
        r.add_runs(q, txt, size=size)
    return r._caption('Figura', caption, r.figs)
